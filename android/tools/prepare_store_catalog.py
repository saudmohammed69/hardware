"""Extract all five store price-list sheets without modifying either workbook.

The final catalogue wins on conflicting prices. Only unambiguous name/size/unit
matches are merged. Different packs, materials, brands and selling bases stay separate.
"""
import collections
import hashlib
import json
import re
import sys
from decimal import Decimal
from fractions import Fraction
from pathlib import Path

import openpyxl
from prepare_final_catalog import number, text

ROOT = Path(__file__).resolve().parents[1]


def clean(value):
    result = text(value)
    return "" if result in {"—", "-"} else result


def name_key(name):
    key = re.sub(r"[^a-z0-9]", "", name.lower())
    # Explicit spelling/word-order equivalents; never strip a brand or material.
    return {"ballvalvepvc": "pvcballvalve", "foltvalve": "floatvalve",
            "floatwall": "floatvalve"}.get(key, key)


def size_key(size):
    value = size.lower().strip().replace('"', ' inch').replace('”', ' inch')
    for symbol, fraction in {"½": "1/2", "¼": "1/4", "¾": "3/4"}.items():
        value = value.replace(symbol, " " + fraction)
    value = re.sub(r"^size\s+", "", value)
    value = re.sub(r"\s+", " ", value).strip()
    mixed = re.fullmatch(r"(\d+) (\d+/\d+)(.*)", value)
    if mixed:
        value = str(float(Fraction(mixed[1]) + Fraction(mixed[2]))) + mixed[3]
    fraction = re.fullmatch(r"(\d+/\d+)(.*)", value)
    if fraction:
        value = str(float(Fraction(fraction[1]))) + fraction[2]
    dimension = re.fullmatch(r"([0-9.]+)\s*(inch|inches|mm|cm|ml|l|kg|g)?", value)
    if dimension:
        amount = Decimal(dimension[1])
        unit = dimension[2] or ""
        if unit == "cm":
            amount, unit = amount * 10, "mm"
        if unit == "l":
            amount, unit = amount * 1000, "ml"
        if unit == "kg":
            amount, unit = amount * 1000, "g"
        if unit == "inches":
            unit = "inch"
        return format(amount.normalize(), "f") + unit
    return re.sub(r"\s+", "", value)


def identity(row):
    return name_key(row["name"]), size_key(row["size"]), row["unit"].lower()


def extract(path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    expected = ["Price List", "Ball Valve", "C-Clip & Allen Key", "Electrical & Fans", "Household Items"]
    assert workbook.sheetnames == expected, workbook.sheetnames
    result = []

    def add(sheet, row, name, size, price, category, kind, unit="Piece", column="C"):
        if not clean(price):
            return
        selling = number(price)
        assert Decimal(selling) > 0
        result.append(dict(name=name, size=clean(size), selling=selling, category=category,
                           type=kind, unit=unit, purchase="", gst="18", stock="", minimum="",
                           sourceSheet=sheet, sourceRow=row, sourceCell=f"{column}{row}"))

    groups = {
        "RAIN & MISC": ("Home Essentials", "Raincoats & Umbrellas"),
        "GARDEN PIPE": ("Plumbing Accessories", "Garden Pipes"),
        "PVC GUM": ("Adhesives & Sealants", "Solvent Cement"),
        "GLASS CUTTER & TOOLS": ("Tools", "Hand Tools"),
        "JOINT MESH": ("Hardware & Fasteners", "Joint Mesh"),
        "TERMINATOR": ("Wood Care", "Terminator"),
        "LIGHTING": ("Electrical", "Lighting"),
        "VALVES": ("Valves", "Valves"),
        "FLEXIBLE PIPE": ("Plumbing Accessories", "Flexible Pipes"),
        "MISCELLANEOUS": ("Home Essentials", "General"),
    }
    section = None
    for i, (name, size, price) in enumerate(workbook["Price List"].values, 1):
        if i < 3:
            continue
        if price is None:
            section = name
            continue
        category, kind = ("Adhesives & Sealants", "Adhesives") if section.startswith("ADHESIVES") else groups[section]
        unit = "Piece"
        if section == "GARDEN PIPE":
            name, diameter = name.rsplit(" ", 1)
            unit = "Roll" if size == "Full roll" else "Meter"
            size = diameter
        elif size == "Per meter":
            unit, size = "Meter", ""
        if name == "Foot Valve":
            kind = "Foot Valves"
        elif name == "Folt Valve":
            name, kind = "Float Valve", "Float Valves"
        add("Price List", i, name, size, price, category, kind, unit)

    for i, values in enumerate(workbook["Ball Valve"].values, 1):
        if i < 3:
            continue
        for j, material in enumerate(("PVC", "UPVC", "CPVC"), 1):
            add("Ball Valve", i, material + " Ball Valve", values[0], values[j], "Valves", "Ball Valve", column=chr(65+j))

    for i, values in enumerate(workbook["C-Clip & Allen Key"].values, 1):
        if i < 3:
            continue
        if values[0]:
            add("C-Clip & Allen Key", i, "C-Clip", values[0], values[1], "Electrical", "Cable Clips", "Pack", "B")
        if values[3]:
            add("C-Clip & Allen Key", i, "Allen Key", values[3], values[4], "Tools", "Allen Keys", column="E")

    section = None
    for i, (name, size, price) in enumerate(workbook["Electrical & Fans"].values, 1):
        if i < 3:
            continue
        if price is None:
            section = name
            continue
        if section.endswith("FAN"):
            kind = section.title() + "s"
            name = name + " " + section.title()
            category = "Electrical"
            if size == "150 mm / 15 cm":
                size = "150 mm"
        else:
            category, kind = "Tools", "Hand Tools"
        add("Electrical & Fans", i, name, size, price, category, kind)

    section = None
    household = {
        "BHATTI (STOVE)": ("Welding Supplies", "Equipment"),
        "CHATTI": ("Home Essentials", "Chatti"),
        "MIRROR": ("Home Essentials", "Mirrors"),
        "DUSTBIN": ("Home Essentials", "Dustbins"),
        "BUCKET": ("Home Essentials", "Buckets & Containers"),
        "MUG & RACK": ("Home Essentials", "Mugs & Racks"),
        "CLEANING TOOLS": ("Cleaning Tools", "Cleaning Tools"),
        "GARDEN / DIGGING TOOLS": ("Tools", "Garden Tools"),
        "GARBAGE & STORAGE": ("Home Essentials", "Garbage Bags"),
        "SEAT COVER": ("Plumbing Accessories", "Seat Covers"),
    }
    for i, (name, size, price) in enumerate(workbook["Household Items"].values, 1):
        if i < 3:
            continue
        if price is None:
            section = name
            continue
        category, kind = household[section]
        if section == "SEAT COVER":
            name += " Seat Cover"
        if name == "Mug":
            category, kind = "Kitchen & Dining", "Mugs"
        if name in {"Mop", "Wiper"}:
            kind = name + "s"
        if clean(size).startswith("Size "):
            size = size.removeprefix("Size ")
        add("Household Items", i, name, size, price, category, kind)
    return result


def prepare(path):
    base = json.loads((ROOT / "app/src/main/assets/final_catalog.json").read_text(encoding="utf-8"))
    by_identity = collections.defaultdict(list)
    for row in base:
        by_identity[identity(row)].append(row)
    extracted = extract(path)
    output = []
    report = []
    for row in extracted:
        matches = by_identity[identity(row)]
        assert len(matches) <= 1, (row, matches)
        if matches:
            master = matches[0]
            target = dict(master)
            target["selling"] = master["selling"] or row["selling"]
            action = "keep_final_price" if master["selling"] else "fill_missing_price"
            target["baseSelling"] = master["selling"]
        else:
            target = dict(row)
            key = "\x1f".join(identity(row))
            target["id"] = "final-store-" + hashlib.sha256(key.encode()).hexdigest()[:32]
            action = "add_product"
        target["storeSource"] = {k: row[k] for k in ("sourceSheet", "sourceRow", "sourceCell", "selling")}
        output.append(target)
        report.append(dict(sheet=row["sourceSheet"], cell=row["sourceCell"], name=row["name"], size=row["size"],
                           unit=row["unit"], storePrice=row["selling"], price=target["selling"], action=action,
                           targetId=target["id"], finalSourceIds=target.get("sourceIds", [])))
    assert len({row["id"] for row in output}) == len(output), "Duplicate supplement rows"
    (ROOT / "app/src/main/assets/store_catalog.json").write_text(json.dumps(output, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    (ROOT / "tools/store_catalog_merge_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(json.dumps(dict(sheets=dict(collections.Counter(r["sheet"] for r in report)),
                         actions=dict(collections.Counter(r["action"] for r in report)))))
    for r in report:
        print(r["action"], r["sheet"], r["cell"], r["name"], r["size"], r["unit"], r["storePrice"], "=>", r["price"])


if __name__ == "__main__":
    prepare(Path(sys.argv[1]))
