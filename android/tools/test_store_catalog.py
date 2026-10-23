"""Checks extraction coverage, price precedence, and pack/size boundaries."""
import collections
import json
import unittest
from pathlib import Path
import openpyxl

from prepare_store_catalog import identity, size_key, ROOT
from prepare_final_catalog import number


class StoreCatalogTest(unittest.TestCase):
    def setUp(self):
        self.rows = json.loads((ROOT / 'app/src/main/assets/store_catalog.json').read_text(encoding='utf-8'))
        self.master = json.loads((ROOT / 'app/src/main/assets/final_catalog.json').read_text(encoding='utf-8'))
        self.by_id = {r['id']: r for r in self.master}

    def test_every_priced_cell_from_every_sheet_is_accounted_for(self):
        # Independent raw-cell inventory, not the extraction routine under test.
        workbook = openpyxl.load_workbook(Path(r'C:\Users\mohammed saud\OneDrive\Documents\store_price_list.xlsx'), data_only=True)
        price_columns = {'Price List': (3,), 'Ball Valve': (2, 3, 4),
                         'C-Clip & Allen Key': (2, 5), 'Electrical & Fans': (3,), 'Household Items': (3,)}
        expected = {}
        for sheet_name, columns in price_columns.items():
            sheet = workbook[sheet_name]
            for row in range(3, sheet.max_row + 1):
                for column in columns:
                    cell = sheet.cell(row, column)
                    if isinstance(cell.value, (int, float)) and cell.value > 0:
                        expected[(sheet_name, cell.coordinate)] = number(cell.value)
        actual = {(r['storeSource']['sourceSheet'], r['storeSource']['sourceCell']): r['storeSource']['selling'] for r in self.rows}
        self.assertEqual(expected, actual)
        self.assertEqual({'Price List': 75, 'Ball Valve': 15, 'C-Clip & Allen Key': 15,
                          'Electrical & Fans': 20, 'Household Items': 81},
                         dict(collections.Counter(r['storeSource']['sourceSheet'] for r in self.rows)))

    def test_matched_master_prices_win_and_only_missing_prices_are_filled(self):
        added = filled = matched = 0
        for row in self.rows:
            base = self.by_id.get(row['id'])
            if base is None:
                added += 1
                self.assertEqual(row['storeSource']['selling'], row['selling'])
            elif base['selling']:
                matched += 1
                self.assertEqual(base['selling'], row['selling'])
            else:
                filled += 1
                self.assertEqual(row['storeSource']['selling'], row['selling'])
        self.assertEqual((173, 3, 30), (added, filled, matched))
        self.assertEqual(len(self.rows), len({r['id'] for r in self.rows}))

    def test_fraction_and_capacity_equivalence_does_not_erase_selling_unit(self):
        self.assertEqual(size_key('1½"'), size_key('1.5 inch'))
        self.assertEqual(size_key('½ L'), size_key('500 ml'))
        self.assertEqual(size_key('Size 01'), size_key('01'))
        self.assertNotEqual(identity(dict(name='Garden Pipe', size='½"', unit='Roll')),
                            identity(dict(name='Garden Pipe', size='½"', unit='Meter')))
        self.assertNotEqual(size_key('Type 1'), size_key('1'))


if __name__ == '__main__':
    unittest.main()
